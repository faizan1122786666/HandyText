from docx import Document
from docx.shared import Pt, RGBColor, Mm, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.shared import OxmlElement, qn
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_RIGHT, TA_JUSTIFY
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfbase.ttfonts import TTFont
from bs4 import BeautifulSoup
import arabic_reshaper
from bidi.algorithm import get_display
import os
import uuid
import re
from xml.sax.saxutils import escape
from .ocr_layout import should_preserve_line_breaks, _TERMINAL_PUNCT


# --- Right-to-left (Arabic-script) support ------------------------------------
# reportlab's built-in fonts (Times-Roman, etc.) contain no Arabic glyphs, so
# RTL text renders as empty boxes. We bundle a Unicode font with Arabic glyphs
# and reshape + reorder the text ourselves because reportlab has no shaping or
# bidirectional layout engine.
_FONT_DIR = os.path.join(os.path.dirname(__file__), "..", "assets", "fonts")
ARABIC_FONT_NAME = "NotoNaskhArabic"
_arabic_font_registered = False

# Arabic, Arabic Supplement, Arabic Extended-A, and the presentation-form blocks.
_ARABIC_RANGES = (
    (0x0600, 0x06FF),
    (0x0750, 0x077F),
    (0x08A0, 0x08FF),
    (0xFB50, 0xFDFF),
    (0xFE70, 0xFEFF),
)


def _register_arabic_font() -> bool:
    """Register the bundled Arabic-capable font with reportlab (once)."""
    global _arabic_font_registered
    if _arabic_font_registered:
        return True
    font_path = os.path.join(_FONT_DIR, "NotoNaskhArabic-Regular.ttf")
    if os.path.exists(font_path):
        try:
            pdfmetrics.registerFont(TTFont(ARABIC_FONT_NAME, font_path))
            _arabic_font_registered = True
        except Exception:
            _arabic_font_registered = False
    return _arabic_font_registered


def _has_arabic(text: str) -> bool:
    for ch in text or "":
        cp = ord(ch)
        for lo, hi in _ARABIC_RANGES:
            if lo <= cp <= hi:
                return True
    return False


def _shape_rtl(text: str) -> str:
    """Connect Arabic-script letters into their contextual forms and apply
    bidirectional reordering so the (shaping-less) PDF renderer lays them out
    correctly."""
    return get_display(arabic_reshaper.reshape(text))


def _wrap_rtl_lines(text: str, font_name: str, font_size: float, max_width: float) -> list[str]:
    """Greedily word-wrap RTL text in *logical* order so multi-line paragraphs
    keep the correct line order. reportlab can't wrap pre-reordered bidi text,
    so we break into fitting lines here and shape each line separately."""
    words = (text or "").split()
    if not words:
        return [text or ""]
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if not current or stringWidth(_shape_rtl(trial), font_name, font_size) <= max_width:
            current = trial
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def _safe_export_name(filename: str) -> str:
    base = os.path.basename(filename or "document")
    base = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", base).strip(" .")
    return base or "document"


def generate_txt(text: str, filename: str) -> str:
    """Generate a .txt file and return its path."""
    file_path = f"exports/{uuid.uuid4()}_{_safe_export_name(filename)}.txt"
    os.makedirs("exports", exist_ok=True)
    with open(file_path, "w", encoding="utf-8") as f:
        f.write(text)
    return file_path


def _css_value(style: str, name: str) -> str:
    match = re.search(rf"{re.escape(name)}\s*:\s*([^;]+)", style or "", flags=re.I)
    return match.group(1).strip().lower() if match else ""


def _is_highlighted(node) -> bool:
    if getattr(node, "name", None) == "mark":
        return True
    style = node.get("style", "") if hasattr(node, "get") else ""
    return bool(_css_value(style, "background-color") or _css_value(style, "background"))


def _is_bold(node) -> bool:
    if getattr(node, "name", None) in ("b", "strong"):
        return True
    style = node.get("style", "") if hasattr(node, "get") else ""
    weight = _css_value(style, "font-weight")
    return weight in ("bold", "bolder", "600", "700", "800", "900")


def _is_italic(node) -> bool:
    if getattr(node, "name", None) in ("i", "em"):
        return True
    return _css_value(node.get("style", "") if hasattr(node, "get") else "", "font-style") == "italic"


def _is_underline(node) -> bool:
    if getattr(node, "name", None) == "u":
        return True
    decoration = _css_value(node.get("style", "") if hasattr(node, "get") else "", "text-decoration")
    return "underline" in decoration


def _is_strike(node) -> bool:
    if getattr(node, "name", None) in ("s", "strike", "del"):
        return True
    decoration = _css_value(node.get("style", "") if hasattr(node, "get") else "", "text-decoration")
    return "line-through" in decoration


def _block_elements(soup: BeautifulSoup):
    blocks = soup.find_all(["p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6"], recursive=True)
    if blocks:
        return blocks
    return [soup]


def _blocks_plain_text(blocks) -> str:
    return "\n".join(element.get_text(" ", strip=True) for element in blocks)


def _is_blank_block(element) -> bool:
    """A block carrying no visible text (e.g. an empty line <div><br></div>)."""
    return not element.get_text(strip=True)


def _is_blank_group(group) -> bool:
    """A paragraph group that is only blank line(s) — rendered as empty lines."""
    return bool(group) and all(_is_blank_block(element) for element in group)


def _is_heading_block(element) -> bool:
    if getattr(element, "name", None) in ("h1", "h2", "h3", "h4", "h5", "h6"):
        return True
    return hasattr(element, "get") and element.get("data-auto-heading") == "true"


# A line that opens with a bullet/number marker is a list item and must not be
# merged into the previous paragraph.
_BULLET_RE = re.compile(r"^\s*([•●▪‣◦*–—-]|\d+[.)])\s+")


def _starts_with_bullet(element) -> bool:
    return bool(_BULLET_RE.match(element.get_text() or ""))


def _is_fully_bold_block(element) -> bool:
    """True when the whole block is bold (a label/heading like 'Functions:')."""
    text = element.get_text(strip=True)
    if not text:
        return False
    bold_text = "".join(b.get_text() for b in element.find_all(["b", "strong"]))
    return bold_text.strip() == text


def _group_blocks_into_paragraphs(blocks):
    """Reflow per-line blocks into flowing paragraphs. The editor stores every
    OCR line as its own <div>, which would otherwise export as a separate
    half-filled line. Soft-wrapped lines (where the previous line ends
    mid-sentence) are merged into one paragraph so text fills the full page
    width, while structural lines stay on their own: a blank line is kept as an
    empty paragraph; headings and fully-bold labels stand alone; each
    bullet/numbered item starts a new paragraph; and a line that follows one
    ending in sentence punctuation begins a new paragraph (so deliberately
    separate lines are not collapsed together onto fewer pages)."""
    if should_preserve_line_breaks(_blocks_plain_text(blocks)):
        # Preserve every line, blank lines included, so the user's spacing and
        # page layout are reproduced exactly.
        return [[element] for element in blocks]

    groups = []
    current = []
    for element in blocks:
        if _is_blank_block(element):
            if current:
                groups.append(current)
                current = []
            # Keep each blank line as its own (empty) paragraph. The user often
            # adds blank lines to push text onto a new page; dropping them
            # collapsed multi-page documents back onto a single page on export.
            groups.append([element])
        elif _is_heading_block(element) or _is_fully_bold_block(element):
            if current:
                groups.append(current)
                current = []
            groups.append([element])
        elif _starts_with_bullet(element):
            if current:
                groups.append(current)
                current = []
            current = [element]
        else:
            prev_text = current[-1].get_text(strip=True) if current else ""
            if current and prev_text.endswith(_TERMINAL_PUNCT):
                # The previous line ends a sentence/label, so this line is a
                # deliberate new line rather than a soft wrap — keep it as its
                # own paragraph instead of merging it onto the previous one.
                groups.append(current)
                current = [element]
            else:
                current.append(element)
    if current:
        groups.append(current)
    return groups


def _reflow_plain_text(text: str) -> list[str]:
    """Reflow raw text into paragraphs: blank lines separate paragraphs and
    single line breaks are treated as soft wraps (joined), so each paragraph
    fills the page width instead of breaking after every short line."""
    if should_preserve_line_breaks(text):
        return [line.strip() for line in (text or "").replace("\r\n", "\n").split("\n") if line.strip()]

    paragraphs = []
    for block in re.split(r"\n[ \t]*\n", (text or "").replace("\r\n", "\n")):
        current = ""
        for line in block.split("\n"):
            stripped = line.strip()
            if not stripped:
                continue
            if _BULLET_RE.match(stripped):
                # A bullet/numbered line starts its own paragraph.
                if current:
                    paragraphs.append(current)
                current = stripped
            else:
                current = stripped if not current else f"{current} {stripped}"
        if current:
            paragraphs.append(current)
    return paragraphs


def _set_paragraph_rtl(paragraph):
    """Mark a Word paragraph as right-to-left and justified (fill width). Word
    performs its own Arabic shaping, so this is all that's needed for correct
    RTL text."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)
    paragraph.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    for run in paragraph.runs:
        rPr = run._r.get_or_add_rPr()
        rtl = OxmlElement('w:rtl')
        rtl.set(qn('w:val'), '1')
        rPr.append(rtl)


def parse_html_to_docx(html_content: str, doc: Document):
    """Parse HTML and add to Word document with formatting."""
    if not html_content:
        return
    
    soup = BeautifulSoup(html_content, 'html.parser')

    def add_runs(paragraph, node, state):
        if getattr(node, "name", None) is None:
            text = str(node).replace("\xa0", " ")
            if text:
                run = paragraph.add_run(text)
                run.bold = state["bold"]
                run.italic = state["italic"]
                run.underline = state["underline"]
                run.font.strike = state["strike"]
                if state["highlight"]:
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            return

        if node.name == "br":
            paragraph.add_run().add_break()
            return

        next_state = {
            "bold": state["bold"] or _is_bold(node),
            "italic": state["italic"] or _is_italic(node),
            "underline": state["underline"] or _is_underline(node),
            "strike": state["strike"] or _is_strike(node),
            "highlight": state["highlight"] or _is_highlighted(node),
        }

        for child in node.children:
            add_runs(paragraph, child, next_state)

    for group in _group_blocks_into_paragraphs(_block_elements(soup)):
        if _is_blank_group(group):
            # Empty paragraph -> one blank line, preserving the user's spacing.
            blank = doc.add_paragraph()
            blank.paragraph_format.line_spacing = 1.15
            continue

        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.15

        is_heading = len(group) == 1 and _is_heading_block(group[0])
        if is_heading:
            paragraph.paragraph_format.space_after = Pt(8)
            base_state = {"bold": True, "italic": False, "underline": False, "strike": False, "highlight": False}
        else:
            base_state = {"bold": False, "italic": False, "underline": False, "strike": False, "highlight": False}

        # Merge the lines of this paragraph, separating them with a space so the
        # text reflows to fill the line rather than breaking after each line.
        for idx, element in enumerate(group):
            if idx:
                paragraph.add_run(" ")
            add_runs(paragraph, element, base_state)

        group_text = " ".join(el.get_text() for el in group)
        if _has_arabic(group_text):
            _set_paragraph_rtl(paragraph)
        elif _BULLET_RE.match(group_text):
            # Hanging indent so wrapped lines align under the text.
            paragraph.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        elif is_heading:
            paragraph.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        else:
            # Justify body text so each line fills the full width.
            paragraph.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

# Map a UI border style to Word's border value + size (size is in eighths of a
# point, matching the editor's solid/thick/double/dashed/dotted options).
_DOCX_BORDER = {
    "solid": ("single", 4),
    "thick": ("single", 24),
    "double": ("double", 6),
    "dashed": ("dashed", 4),
    "dotted": ("dotted", 4),
}


def generate_docx(html_content: str, plain_text: str, filename: str,
                  page_border: bool = True, page_border_style: str = "solid") -> str:
    """Generate a .docx file and return its path."""
    file_path = f"exports/{uuid.uuid4()}_{_safe_export_name(filename)}.docx"
    os.makedirs("exports", exist_ok=True)
    doc = Document()

    # Document title metadata so viewers show the user's name, not "(anonymous)".
    doc_title = _safe_export_name(os.path.splitext(os.path.basename(filename or "document"))[0])
    doc.core_properties.title = doc_title
    doc.core_properties.author = "HandyText"

    # Set page size to A4 (210mm x 297mm)
    section = doc.sections[0]
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.left_margin = Pt(72)      # 1 inch
    section.right_margin = Pt(72)     # 1 inch
    section.top_margin = Pt(72)       # 1 inch
    section.bottom_margin = Pt(72)    # 1 inch
    
    # Set page border
    def set_page_border(section, border_val='single', border_size=4, border_color='000000'):
        p = section._sectPr
        pgBorders = OxmlElement('w:pgBorders')
        pgBorders.set(qn('w:offsetFrom'), 'page')

        for border in ['top', 'left', 'bottom', 'right']:
            bd = OxmlElement(f'w:{border}')
            bd.set(qn('w:val'), border_val)
            bd.set(qn('w:sz'), str(border_size))
            bd.set(qn('w:space'), '24')
            bd.set(qn('w:color'), border_color)
            pgBorders.append(bd)

        p.append(pgBorders)

    border_style = (page_border_style or "solid").lower()
    if page_border and border_style != "none":
        val, size = _DOCX_BORDER.get(border_style, _DOCX_BORDER["solid"])
        set_page_border(section, border_val=val, border_size=size)

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.size = Pt(12)
    font.name = 'Times New Roman'
    
    if html_content:
        parse_html_to_docx(html_content, doc)
    else:
        for para_text in _reflow_plain_text(plain_text):
            paragraph = doc.add_paragraph()
            paragraph.add_run(para_text)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.15
            if _has_arabic(para_text):
                _set_paragraph_rtl(paragraph)
            elif _BULLET_RE.match(para_text):
                paragraph.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
                paragraph.paragraph_format.left_indent = Inches(0.3)
                paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            else:
                paragraph.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    
    doc.save(file_path)
    return file_path


def _html_to_reportlab_markup(html_content: str) -> list[tuple[str, str]]:
    """Return a list of (reportlab_markup, plain_text) tuples, one per block.
    The plain text is used to decide whether a paragraph needs RTL handling."""
    soup = BeautifulSoup(html_content, "html.parser")

    def node_to_markup(node, state) -> str:
        if getattr(node, "name", None) is None:
            return escape(str(node).replace("\xa0", " "))
        if node.name == "br":
            return "<br/>"

        next_state = {
            "bold": state["bold"] or _is_bold(node),
            "italic": state["italic"] or _is_italic(node),
            "underline": state["underline"] or _is_underline(node),
            "strike": state["strike"] or _is_strike(node),
            "highlight": state["highlight"] or _is_highlighted(node),
        }

        inner = "".join(node_to_markup(child, next_state) for child in node.children)
        if not inner:
            return ""
        if not state["highlight"] and next_state["highlight"]:
            inner = f'<font backColor="yellow">{inner}</font>'
        if not state["strike"] and next_state["strike"]:
            inner = f"<strike>{inner}</strike>"
        if not state["underline"] and next_state["underline"]:
            inner = f"<u>{inner}</u>"
        if not state["italic"] and next_state["italic"]:
            inner = f"<i>{inner}</i>"
        if not state["bold"] and next_state["bold"]:
            inner = f"<b>{inner}</b>"
        return inner

    paragraphs = []
    for group in _group_blocks_into_paragraphs(_block_elements(soup)):
        if _is_blank_group(group):
            # Blank line -> one empty line (keeps the user's spacing/pagination).
            paragraphs.append(("&nbsp;", ""))
            continue
        is_heading = len(group) == 1 and _is_heading_block(group[0])
        parts = []
        for element in group:
            base_state = {
                "bold": is_heading,
                "italic": False,
                "underline": False,
                "strike": False,
                "highlight": False,
            }
            content = "".join(node_to_markup(child, base_state) for child in element.children).strip()
            if content:
                parts.append(content)
        # Join the merged lines with a space so the paragraph fills the width.
        markup = " ".join(parts)
        if is_heading and markup:
            markup = f"<b>{markup}</b>"
        plain = " ".join(el.get_text() for el in group)
        paragraphs.append((markup or "&nbsp;", plain))
    return paragraphs

def generate_pdf(html_content: str, plain_text: str, filename: str,
                 page_border: bool = True, page_border_style: str = "solid") -> str:
    """Generate a .pdf file and return its path."""
    file_path = f"exports/{uuid.uuid4()}_{_safe_export_name(filename)}.pdf"
    os.makedirs("exports", exist_ok=True)

    # Document title metadata so PDF viewers show the user's name, not "(anonymous)".
    doc_title = _safe_export_name(os.path.splitext(os.path.basename(filename or "document"))[0])

    doc = SimpleDocTemplate(
        file_path,
        pagesize=A4,
        leftMargin=1*inch,
        rightMargin=1*inch,
        topMargin=1*inch,
        bottomMargin=1*inch,
        title=doc_title,
        author="HandyText",
    )
    
    styles = getSampleStyleSheet()

    arabic_ok = _register_arabic_font()

    # Body style: justified so each line fills the full width edge-to-edge
    # (the last/only line of a paragraph stays left, the default justify
    # behaviour, so short lines and headings are not stretched).
    normal_style = ParagraphStyle(
        'CustomNormal',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=12,
        leading=15,
        alignment=TA_JUSTIFY,
        leftIndent=0,
        rightIndent=0,
        firstLineIndent=0,
        spaceAfter=4
    )

    # Bullet/numbered list items: left-aligned with a hanging indent so wrapped
    # continuation lines align under the text, not under the marker.
    bullet_style = ParagraphStyle(
        'CustomBullet',
        parent=normal_style,
        alignment=TA_LEFT,
        leftIndent=20,
        bulletIndent=6,
        spaceAfter=4,
    )

    # Smart fill-width for RTL / Arabic-script lines: lines long enough to wrap
    # are justified (fill width); short single lines stay at the right margin
    # (natural RTL), avoiding the gappy look of stretching a short line.
    arabic_font = ARABIC_FONT_NAME if arabic_ok else 'Times-Roman'
    # Last/short physical line of a paragraph: sit at the right margin.
    arabic_right = ParagraphStyle(
        'ArabicRight', parent=normal_style, fontName=arabic_font, alignment=TA_RIGHT,
    )
    # A "full" physical line that should stretch edge-to-edge. justifyLastLine=1
    # forces the (single-line) flowable to justify; spaceAfter=0 keeps wrapped
    # lines of one paragraph tight together.
    arabic_justify = ParagraphStyle(
        'ArabicJustify', parent=normal_style, fontName=arabic_font,
        alignment=TA_JUSTIFY, justifyLastLine=1, spaceAfter=0,
    )
    # Text width available inside the 1-inch left/right margins, minus the
    # reportlab Frame's default 6pt padding on each side (plus a small safety
    # margin) so our pre-wrapped lines are never re-wrapped by reportlab.
    avail_width = A4[0] - 2 * inch - 12 - 2

    story = []

    if html_content:
        paragraphs = _html_to_reportlab_markup(html_content)
    else:
        paragraphs = [(escape(p), p) for p in _reflow_plain_text(plain_text)]

    for markup, plain in paragraphs:
        if not (markup and markup.strip()):
            story.append(Spacer(1, 6))
        elif arabic_ok and _has_arabic(plain):
            # Wrap in logical order, then shape each physical line. Full lines
            # fill the width; the final (short) line sits at the right margin.
            # Inline markup is dropped for RTL lines so shaping gets a clean run.
            lines = _wrap_rtl_lines(plain, arabic_font, normal_style.fontSize, avail_width)
            for i, line in enumerate(lines):
                is_last = i == len(lines) - 1
                story.append(Paragraph(escape(_shape_rtl(line)),
                                       arabic_right if is_last else arabic_justify))
        else:
            bullet_match = _BULLET_RE.match(plain)
            if bullet_match:
                marker = bullet_match.group(1)
                bullet = marker if marker[0].isdigit() else "•"
                body = _BULLET_RE.sub("", markup, count=1)
                story.append(Paragraph(body or "&nbsp;", bullet_style, bulletText=bullet))
            else:
                story.append(Paragraph(markup, normal_style))

    # Custom page template with an optional border: a rectangle inset from the
    # page edge with connected corners (matches the editor view). The line style
    # mirrors the editor's solid/thick/double/dashed/dotted options.
    border_style = (page_border_style or "solid").lower()

    def on_page(canvas, doc):
        if not page_border or border_style == "none":
            return
        canvas.saveState()
        canvas.setStrokeColorRGB(0, 0, 0)
        canvas.setLineWidth(4 if border_style == "thick" else 2)
        if border_style == "dashed":
            canvas.setDash(6, 3)
        elif border_style == "dotted":
            canvas.setDash(1, 3)
        m = 0.5 * inch       # inset margin from the page edge
        page_width, page_height = A4
        canvas.rect(m, m, page_width - 2 * m, page_height - 2 * m, stroke=1, fill=0)
        if border_style == "double":
            # A second inner rectangle gives the classic double-line look.
            g = 3
            canvas.rect(m + g, m + g, page_width - 2 * m - 2 * g,
                        page_height - 2 * m - 2 * g, stroke=1, fill=0)
        canvas.restoreState()

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    
    return file_path
